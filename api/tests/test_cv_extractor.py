"""Tests for CV file text extraction."""

import asyncio
from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from app.services.cv_extractor import CVExtractor


def test_extracts_text_from_docx():
    """DOCX paragraphs should be converted to plain text."""

    document = Document()
    document.add_paragraph("Jane Doe")
    document.add_paragraph("Professional Summary")
    document.add_paragraph("Customer service specialist.")

    buffer = BytesIO()
    document.save(buffer)

    result = asyncio.run(CVExtractor().extract(
        filename="cv.docx",
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        file_bytes=buffer.getvalue(),
    ))

    assert "Jane Doe" in result
    assert "Professional Summary" in result
    assert "Customer service specialist." in result


def test_extracts_text_from_pdf():
    """PDF files should be accepted and processed."""

    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)

    buffer = BytesIO()
    writer.write(buffer)

    result = asyncio.run(CVExtractor().extract(
        filename="cv.pdf",
        content_type="application/pdf",
        file_bytes=buffer.getvalue(),
    ))

    assert result == ""


def test_rejects_unsupported_file_type():
    """Unsupported CV formats should raise a clear error."""

    with pytest.raises(ValueError, match="Unsupported CV format"):
        asyncio.run(CVExtractor().extract(
            filename="cv.txt",
            content_type="text/plain",
            file_bytes=b"Jane Doe",
        ))


def test_rejects_empty_file():
    """Empty uploads should raise a clear error."""

    with pytest.raises(ValueError, match="uploaded CV is empty"):
        asyncio.run(CVExtractor().extract(
            filename="cv.pdf",
            content_type="application/pdf",
            file_bytes=b"",
        ))


def test_extracts_text_from_docx_tables():
    """DOCX table content should be included in extracted text."""

    document = Document()
    table = document.add_table(rows=2, cols=2)

    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "Python, FastAPI"
    table.cell(1, 0).text = "Experience"
    table.cell(1, 1).text = "Backend Developer"

    buffer = BytesIO()
    document.save(buffer)

    result = asyncio.run(
        CVExtractor().extract(
            filename="table_cv.docx",
            content_type=(
                "application/vnd.openxmlformats-officedocument."
                "wordprocessingml.document"
            ),
            file_bytes=buffer.getvalue(),
        )
    )

    assert "Skills | Python, FastAPI" in result
    assert "Experience | Backend Developer" in result
