"""CV file text extraction service for JobMatch."""

import re
from io import BytesIO

from docx import Document
import pdfplumber


class CVExtractor:
    # pylint: disable=too-few-public-methods
    """Extract plain text from supported CV files."""

    SUPPORTED_TYPES = {
        "application/pdf": "pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    }

    async def extract(
        self,
        filename: str,
        content_type: str,
        file_bytes: bytes,
    ) -> str:
        """Extract text from a PDF or DOCX CV."""

        del filename

        file_type = self.SUPPORTED_TYPES.get(content_type)

        if not file_type:
            raise ValueError(
                "Unsupported CV format. Please upload a PDF or DOCX file."
            )

        if not file_bytes:
            raise ValueError("The uploaded CV is empty.")

        if file_type == "pdf":
            return self._extract_pdf(file_bytes)

        return self._extract_docx(file_bytes)

    @staticmethod
    def _extract_pdf(file_bytes: bytes) -> str:
        pages = []

        with pdfplumber.open(BytesIO(file_bytes)) as pdf:
            for page in pdf.pages:
                text = page.extract_text() or ""

                if text.strip():
                    pages.append(text.strip())

        return CVExtractor._clean_text("\n\n".join(pages))

    @staticmethod
    def _extract_docx(file_bytes: bytes) -> str:
        document = Document(BytesIO(file_bytes))

        paragraphs = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                paragraphs.append(text)

        for table in document.tables:
            for row in table.rows:
                cells = []

                for cell in row.cells:
                    cell_text = cell.text.strip()

                    if cell_text:
                        cells.append(cell_text)

                if cells:
                    paragraphs.append(" | ".join(cells))

        return CVExtractor._clean_text("\n\n".join(paragraphs))

    @staticmethod
    def _clean_text(text: str) -> str:
        """Normalize extracted CV text."""

        text = text.replace("-\n", "-")

        lines = [
            line.strip()
            for line in text.splitlines()
            if line.strip()
        ]

        cleaned = "\n".join(lines)

        cleaned = re.sub(
            r"(?<=[a-z])(?=(application|authentication|automated|RESTAPI))",
            " ",
            cleaned,
            flags=re.IGNORECASE,
        )

        return re.sub(
            r"-\s+",
            "-",
            cleaned,
        )
